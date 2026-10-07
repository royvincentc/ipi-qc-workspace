import tempfile,unittest
from pathlib import Path
from docx import Document
from docx_worker import demo_template,prepare_template,package,roots,text,validate_template,render,NS,W
from environmental_formats import pattern,FAMILIES

class FormatTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.folder=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def source(self,headers,rows):
        file=self.folder/'source.docx';demo_template(file);d=Document(file)
        old=d.tables[0];old._element.getparent().remove(old._element)
        d.add_paragraph('Environmental Monitoring');d.add_paragraph('Name of Sample');d.add_paragraph(':');d.add_paragraph('Confirmed water location')
        t=d.add_table(rows=1,cols=len(headers))
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=value
        d.save(file);return file
    def test_water_four_columns_keeps_independent_result_and_remark_cells(self):
        source=self.source(['ANALYSIS DESIRED:','Standard Specifications','Results','Remarks'],[['Standard Plate Count (SPC)','100 cfu/mL','TNTC','Failed'],['Molds and Yeast','100 cfu/mL','23','Passed']])
        p=pattern(roots(package(source.read_bytes())));self.assertEqual(p['family'],FAMILIES[2]);self.assertEqual(len(p['blocks'][0]['instances']),2)
        self.assertNotIn('TNTC',str(p));self.assertNotIn('Passed',str(p))
        out=self.folder/'prepared.docx';info=prepare_template(source.read_bytes(),out)
        self.assertEqual(info['family'],FAMILIES[2]);self.assertIn('value',validate_template(out.read_bytes())['tokens'])
        self.assertNotIn('Failed',' '.join(text(r) for r in roots(package(out.read_bytes())).values()))
    def test_gip_top_bottom_inherit_criteria_without_importing_results(self):
        source=self.source(['ANALYSIS DESIRED:','Standard Specifications','Results','Remarks'],[['Standard Plate Count (SPC)','Nmt 100 cfu/mL','',''],['Top','','4','Passed'],['Bottom','','5','Passed']])
        p=pattern(roots(package(source.read_bytes())),FAMILIES[1]);rows=p['blocks'][0]['instances']
        self.assertEqual([r['location'] for r in rows],['Top','Bottom']);self.assertEqual([r['criterion'] for r in rows],['Nmt 100 cfu/mL']*2)
        out=self.folder/'gip.docx';info=prepare_template(source.read_bytes(),out,FAMILIES[1]);self.assertEqual(info['blocks']['tests']['mergeColumns'],[])
        self.assertTrue({'test','location','criterion','value','remarks'}<=set(validate_template(out.read_bytes())['tokens']))
        self.assertNotIn('Passed',' '.join(text(r) for r in roots(package(out.read_bytes())).values()))
    def test_warehouse_preserves_both_channels_and_phase(self):
        source=self.source(['ANALYSIS DESIRED:','Phase','Area','Standard Specifications','Results','','Remarks'],[['','','','','Active Air','Passive Air',''],['Standard Plate Count (SPC)','1','Center','-----','251 cfu/m3','27 cfu','Passed']])
        p=pattern(roots(package(source.read_bytes())));rows=p['blocks'][0]['instances']
        self.assertEqual(p['family'],FAMILIES[3]);self.assertEqual([r['channel'] for r in rows],['active-air','passive-air']);self.assertTrue(all(r['unit']=='' for r in rows));self.assertNotIn('251',str(p))
        out=self.folder/'warehouse.docx';prepare_template(source.read_bytes(),out)
        tokens=validate_template(out.read_bytes())['tokens'];self.assertTrue({'phase','activeValue','passiveValue'}<=set(tokens))
        fields={t:'' for t in tokens};report=self.folder/'generated.docx'
        render(out,{'fields':fields,'rows':[{'test':'SPC','phase':'1','location':'Center','criterion':'Owner-confirmed criterion','activeValue':'0 cfu/m3','passiveValue':'2 cfu','remarks':'Active: Passed\nPassive: Failed'}]},report)
        all_text=' '.join(text(r) for r in roots(package(report.read_bytes())).values());self.assertIn('0 cfu/m3',all_text);self.assertIn('2 cfu',all_text);self.assertNotIn('251',all_text)

    def test_master_layout_expands_named_blocks_without_losing_or_copying_values(self):
        source=self.source(['ANALYSIS DESIRED:','Area','Standard Specifications','Results','Remarks'],[['Standard Plate Count (SPC)','Nozzle','Nmt 100 cfu','41','Passed']])
        master=self.folder/'master.docx';prepare_template(source.read_bytes(),master)
        tokens=validate_template(master.read_bytes())['tokens'];fields={t:'' for t in tokens}
        output=self.folder/'multi.docx'
        row={'test':'SPC','location':'Nozzle','criterion':'Nmt 100 cfu','value':'0','remarks':'','groupKey':'spc'}
        render(master,{'fields':fields,'blocks':{'spc':[row],'my':[{**row,'test':'MY','value':'2','groupKey':'my'}]},'renderOptions':{'adaptiveBlocks':True,'rowGrouping':{'mergeColumns':[0,2]}}},output)
        document=Document(output);tables=[t for t in document.tables if 'ANALYSIS DESIRED:' in t.rows[0].cells[0].text]
        self.assertEqual(len(tables),2);self.assertEqual(tables[0].rows[1].cells[3].text,'0');self.assertEqual(tables[1].rows[1].cells[3].text,'2')
        self.assertEqual(tables[0]._element.getnext().tag,W+'p')
        self.assertIs(tables[0]._element.getnext().getnext(),tables[1]._element)
        self.assertNotIn('{{',' '.join(text(r) for r in roots(package(output.read_bytes())).values()))

if __name__=='__main__':unittest.main()
