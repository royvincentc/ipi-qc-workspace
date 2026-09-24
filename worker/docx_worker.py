"""Private OOXML inspection and generation. No macros, external links or remote renderer."""
import argparse, copy, hashlib, io, json, re, subprocess, tempfile, zipfile
from pathlib import Path
from lxml import etree as E

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+NS['w']+'}'
XML=E.XMLParser(resolve_entities=False, no_network=True)

def package(data):
    z=zipfile.ZipFile(io.BytesIO(data))
    if sum(x.file_size for x in z.infolist())>40_000_000 or len(z.infolist())>1000:
        raise ValueError('Document package exceeds safety limit')
    parts={n:z.read(n) for n in z.namelist()}
    for n,b in parts.items():
        if 'vbaProject' in n or n.endswith('.bin') or n.startswith('word/embeddings/'):
            raise ValueError('Embedded objects or macros are not accepted')
        if n.endswith('.rels') and b'TargetMode="External"' in b:
            raise ValueError('External document relationships must be removed before import')
    return parts

def roots(parts):
    return {n:E.fromstring(b,XML) for n,b in parts.items() if re.match(r'word/(document|header\d+|footer\d+)\.xml$',n)}

def text(node):
    return ''.join(node.xpath('.//w:t/text()',namespaces=NS))

def set_text(node,value):
    texts=node.xpath('.//w:t',namespaces=NS)
    if texts:
        texts[0].text=value
        for t in texts[1:]:t.text=''
    else:
        p=node if node.tag==W+'p' else E.SubElement(node,W+'p')
        r=E.SubElement(p,W+'r');E.SubElement(r,W+'t').text=value

def set_noted_by(docs,name='{{report.notedBy}}',role='Assistant Head, Microbiology Laboratory'):
    """User-confirmed standing name; signature and signing date remain separate."""
    for root in docs.values():
        for table in root.xpath('.//w:tbl',namespaces=NS):
            if 'Noted by:' not in text(table):continue
            for cell in table.xpath('./w:tr/w:tc',namespaces=NS):
                paragraphs=cell.findall(W+'p')
                for i,p in enumerate(paragraphs):
                    if text(p).strip()=='Assistant Head, Microbiology Laboratory' and i>0:
                        set_text(paragraphs[i-1],name)
                        set_text(p,role)

def prepare_template(data,output):
    """Prepare recognized IPI layouts, never declare them verified automatically."""
    parts=package(data);docs=roots(parts);issues=[];required=[];mapped=[];instances=[]
    labels={'Category':'sample.category','Date&Time Received':'sample.received','Name of Sample':'sample.name','Date&Time Released':'releaseDate','Date Mfd.':'manufactureDate','Batch/Lot No.':'sample.batch','Fill Vol./Wt.':'fillVolume','Expiry Date':'expiryDate','Batch/Lot Size':'batchSize','Requested by':'requestedBy','Purpose':'purpose','Logbook':'logbookReference'}
    found=set()
    for part,root in docs.items():
        paras=root.xpath('.//w:p',namespaces=NS)
        for i,p in enumerate(paras):
            label=text(p).strip()
            if label in labels:
                following=paras[i+1:i+4]
                colon=next((j for j,x in enumerate(following) if text(x).strip()==':'),None)
                if colon is None or colon+1>=len(following):
                    issues.append('Unrecognized field structure: '+label);continue
                value=following[colon+1];key=labels[label]
                set_text(value,'{{'+key+'}}'+(' {{sample.ml}}' if key=='logbookReference' else ''))
                # Legacy reports can store a person's initial as an automatic list marker.
                # It is not part of w:t and would otherwise survive sanitization.
                props=value.find(W+'pPr')
                numbering=props.find(W+'numPr') if props is not None else None
                if numbering is not None:props.remove(numbering)
                found.add(key);mapped.append({'part':part,'field':key})
                if not key.startswith('sample.') and key!='releaseDate':required.append(key)
        if part=='word/document.xml':
            for table in root.xpath('.//w:body/w:tbl',namespaces=NS):
                rows=table.findall(W+'tr')
                if not rows or not any('Results' in text(r) or 'ACTUAL RESULTS' in text(r) for r in rows[:2]):continue
                for ri,row in enumerate(rows):
                    cells=row.findall(W+'tc')
                    if len(cells)<4:continue
                    name,criterion=text(cells[0]),text(cells[1])
                    if not re.search(r'\bNmt\b|\bcfu\b|^Negative$|^Absent',criterion,re.I):continue
                    # Parent labels with subsequent per-container rows do not receive results.
                    next_cells=rows[ri+1].findall(W+'tc') if ri+1<len(rows) else []
                    next_label=text(next_cells[0]).strip() if next_cells else ''
                    is_container=lambda x:bool(re.match(r'^(Bottle\s*\d+|Nozzle\s*\d+|Top(?:-Middle)?|Middle(?:-Bottom)?|Bottom|Tank)$',x,re.I))
                    if not is_container(name.strip()) and is_container(next_label):
                        set_text(cells[2],'');set_text(cells[3],'');continue
                    index=len(instances);instances.append({'index':index,'label':name,'originalCriterion':criterion})
                    set_text(cells[1],'{{result.'+str(index)+'.criterion}}')
                    set_text(cells[2],'{{result.'+str(index)+'.value}}')
                    set_text(cells[3],'{{result.'+str(index)+'.remarks}}')
    # Clear historical sign-off regions, preserving role labels and pagination shapes.
    for part,root in docs.items():
        regions=[root] if 'footer' in part else root.xpath('.//w:tbl[contains(string(.),"Analyzed by:")]',namespaces=NS)
        for region in regions:
            for p in region.xpath('.//w:p',namespaces=NS):
                if p.xpath('.//w:p',namespaces=NS):continue  # preserve nested page-number textboxes
                value=text(p).strip()
                if not value:continue
                if value.startswith('REMARKS:'):set_text(p,'REMARKS: {{overallRemarks}}');continue
                if value.startswith('Date:'):set_text(p,'Date: ');continue
                if re.search(r'^(Page |cc\.|Analyzed by:|Noted by:|By:|APPROVED|REJECTED)',value,re.I):continue
                if re.search(r'Analyst|Laboratory|Head,|Microbiolog',value,re.I):continue
                set_text(p,'')
            for image in region.xpath('.//*[local-name()="imagedata" or local-name()="blip"]'):
                image.getparent().remove(image)
    set_noted_by(docs)
    missing={'sample.name','sample.batch','sample.category','sample.received','logbookReference'}-found
    if missing:issues.append('Unmapped identifiers: '+', '.join(sorted(missing)))
    if not instances:issues.append('No recognized result rows')
    for n,root in docs.items():parts[n]=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    # Core metadata, comments and revisions may also carry historical authorship.
    for n in list(parts):
        if n=='docProps/core.xml':
            r=E.fromstring(parts[n],XML)
            for el in list(r):r.remove(el)
            parts[n]=E.tostring(r,xml_declaration=True,encoding='UTF-8')
    for root in docs.values():
        if root.xpath('.//w:ins|.//w:del|.//w:commentRangeStart',namespaces=NS):issues.append('Tracked changes or comments need manual removal')
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
    return {'requiredFields':sorted(set(required+['analysisDate','logbookReference'])),'mappedFields':mapped,'instances':instances,'issues':issues,'verified':False,'sourceSha256':hashlib.sha256(data).hexdigest(),'instruction':'Review all variable fields, blank signature areas and rendered layout before registration.'}

def inventory(data,name=''):
    parts=package(data); docs=roots(parts); paragraphs=[]; tables=[]; criteria=[]
    for part,root in docs.items():
        for i,p in enumerate(root.xpath('.//w:p',namespaces=NS)):
            if text(p).strip():paragraphs.append({'part':part,'paragraph':i,'text':text(p)})
        for ti,t in enumerate(root.xpath('.//w:tbl',namespaces=NS)):
            rows=[[text(c) for c in r.findall(W+'tc')] for r in t.findall(W+'tr')]
            tables.append({'part':part,'table':ti,'rows':rows})
            for ri,row in enumerate(rows):
                if len(row)>1 and (re.search(r'\bNmt\b|\bcfu\b|^Negative$|^Absent',row[1],re.I)):
                    criteria.append({'testLabel':row[0],'criterion':row[1],'unit':next(iter(re.findall(r'cfu\s*/\s*\w+',row[1],re.I)),''),'sourceLocation':f'{part}:table[{ti}]:row[{ri}]:cell[1]'})
    all_text='\n'.join(p['text'] for p in paragraphs)
    fields={}
    for key,label in [('product','Name of Sample'),('category','Category'),('release','Date&Time Released'),('analysis','Date:'),('batch','Batch/Lot No.')]:
        hits=[p for p in paragraphs if label.lower() in p['text'].lower()]
        fields[key]=hits
    labels=' '.join(c['testLabel'] for c in criteria)
    family='container-location' if re.search(r'Bottle \d|Nozzle|Top|Bottom|Tank',labels) else 'routine'
    if re.search(r'Withdrawal|Stability|Stab\.',all_text,re.I):family='stability'
    if re.search(r'Raw Material|Empty Bottle',all_text,re.I):family='raw-material-packaging'
    return {'name':name,'sha256':hashlib.sha256(data).hexdigest(),'family':family,'paragraphs':paragraphs,'tables':tables,'candidateCriteria':criteria,'fieldEvidence':fields,'issues':['Product, context and explicit report date require confirmation from the evidence.'],'layout':{'sections':[E.tostring(s).decode() for root in docs.values() for s in root.xpath('.//w:sectPr',namespaces=NS)]}}

def replace_tokens(paragraph,values):
    # Preserve each run's original style and replace across split Word runs.
    nodes=paragraph.xpath('.//w:t',namespaces=NS)
    source=''.join(n.text or '' for n in nodes)
    for match in reversed(list(re.finditer(r'\{\{\s*([\w.:-]+)\s*\}\}',source))):
        key=match.group(1)
        if key not in values:raise ValueError('Unknown or missing template field: '+key)
        start,end=match.span(); offset=0; assigned=False
        for node in nodes:
            value=node.text or ''; a,b=offset,offset+len(value);offset=b
            if b<=start or a>=end:continue
            left=value[:max(0,start-a)];right=value[max(0,end-a):]
            node.text=left+(str(values[key]) if not assigned else '')+right
            node.set('{http://www.w3.org/XML/1998/namespace}space','preserve');assigned=True

def render(template,payload,output,soffice=None):
    parts=package(Path(template).read_bytes()); docs=roots(parts)
    fields=payload['fields']; observations=payload['rows']
    report=payload.get('reportSettings')
    if report:set_noted_by(docs,report['notedBy'],report['notedByRole'])
    repeats=0
    for root in docs.values():
        for row in list(root.xpath('.//w:tr',namespaces=NS)):
            if '{{tests}}' not in text(row):continue
            parent=row.getparent();index=parent.index(row);parent.remove(row);repeats+=1
            for i,result in enumerate(observations):
                clone=copy.deepcopy(row)
                for p in clone.xpath('.//w:p',namespaces=NS):replace_tokens(p,{**fields,**result,'tests':''})
                parent.insert(index+i,clone)
        for p in root.xpath('.//w:p',namespaces=NS):replace_tokens(p,fields)
        if '{{' in text(root):raise ValueError('Unresolved template tokens remain')
    if repeats>1:raise ValueError('Only one repeating test block is supported per template; use explicit result tokens for complex layouts')
    for n,root in docs.items():parts[n]=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    settings=E.fromstring(parts['word/settings.xml'],XML) if 'word/settings.xml' in parts else E.Element(W+'settings',nsmap={'w':NS['w']})
    update=settings.find(W+'updateFields')
    if update is None:update=E.SubElement(settings,W+'updateFields')
    update.set(W+'val','true');parts['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8')
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
    if soffice:
        with tempfile.TemporaryDirectory() as tmp:
            profile=Path(tmp,'profile').as_uri()
            subprocess.run([soffice,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',str(Path(output).parent),str(output)],check=True,timeout=90,capture_output=True)
        if not Path(output).with_suffix('.pdf').exists():raise ValueError('Renderer did not produce a PDF preview')
    return {'docx':str(output),'pdf':str(Path(output).with_suffix('.pdf')) if soffice else None}

def validate_template(data):
    parts=package(data);docs=roots(parts);tokens=[]
    for root in docs.values():tokens+=re.findall(r'\{\{\s*([\w.:-]+)\s*\}\}',text(root))
    if not {'sample.name','sample.ml','sample.batch'}.issubset(tokens):raise ValueError('Template needs sample.name, sample.ml and sample.batch placeholders')
    if 'tests' not in tokens and not any(t.startswith('result.') for t in tokens):raise ValueError('Template needs a results block or explicit result placeholders')
    # Signed reference files are never automatically accepted as sanitized templates.
    drawings=sum(len(r.xpath('.//w:drawing|.//w:pict',namespaces=NS)) for r in docs.values())
    fields=' '.join(' '.join(r.xpath('.//w:instrText/text()',namespaces=NS)) for r in docs.values())
    if not re.search(r'\bPAGE\b',fields) or not re.search(r'\bNUMPAGES\b',fields):raise ValueError('Use automatic PAGE and NUMPAGES fields')
    return {'tokens':sorted(set(tokens)),'drawingsRequiringReview':drawings,'requiresHumanSanitizationReview':True}

def demo_template(output):
    from docx import Document
    from docx.shared import Inches,Pt
    d=Document();s=d.sections[0];s.header.paragraphs[0].text='DE-IDENTIFIED DEVELOPMENT FIXTURE · NOT AN IPI APPROVED FORM'
    d.add_heading('Microbiology Analysis Report',0)
    for line in ['Sample: {{sample.name}}','ML number: {{sample.ml}}    Batch: {{sample.batch}}','Date received: {{sample.received}}','Analysis date: {{analysisDate}}    Logbook: {{logbookReference}}']:
        d.add_paragraph(line)
    t=d.add_table(rows=2,cols=4);t.style='Table Grid'
    for c,label in zip(t.rows[0].cells,['Analysis desired','Standard specification','Actual result','Remarks']):c.text=label
    for c,label in zip(t.rows[1].cells,['{{tests}}{{test}}','{{criterion}}','{{value}}','{{remarks}}']):c.text=label
    d.add_paragraph('Analyzed by: {{analyst}}');d.add_paragraph('Reviewed by: __________________    Date: __________')
    p=s.footer.paragraphs[0];p.add_run('Page ')
    for instruction in ['PAGE','NUMPAGES']:
        if instruction=='NUMPAGES':p.add_run(' of ')
        for tag,value in [('fldChar','begin'),('instrText',instruction),('fldChar','end')]:
            r=E.SubElement(p._p,W+'r');el=E.SubElement(r,W+tag)
            if tag=='fldChar':el.set(W+'fldCharType',value)
            else:el.text=value
    d.save(output)

def main():
    p=argparse.ArgumentParser();p.add_argument('action',choices=['inventory','generate','validate','demo','prepare']);p.add_argument('--input');p.add_argument('--payload');p.add_argument('--output');p.add_argument('--soffice');args=p.parse_args()
    if args.action=='inventory':
        path=Path(args.input)
        if path.suffix.lower()=='.zip':
            result=[]
            with zipfile.ZipFile(path) as z:
                for item in z.infolist():
                    if item.filename.endswith('.docx') and not Path(item.filename).name.startswith('~$'):
                        try:result.append(inventory(z.read(item),item.filename))
                        except Exception as e:result.append({'name':item.filename,'error':str(e)})
        else:result=inventory(path.read_bytes(),path.name)
    elif args.action=='generate':result=render(args.input,json.loads(Path(args.payload).read_text(encoding='utf8')),args.output,args.soffice)
    elif args.action=='validate':result=validate_template(Path(args.input).read_bytes())
    elif args.action=='prepare':result=prepare_template(Path(args.input).read_bytes(),args.output)
    else:demo_template(args.output);result={'path':args.output}
    print(json.dumps(result,ensure_ascii=True))
if __name__=='__main__':main()
