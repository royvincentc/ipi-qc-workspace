import tempfile,unittest,zipfile,json
from pathlib import Path
from docx import Document
from environmental_batch import scan

class BatchTests(unittest.TestCase):
    def report(self,folder,category,value):
        file=folder/'report.docx';d=Document()
        for label,text in [('Category',category),('Name of Sample','Boost'),('Area','Compounding'),('Batch/Lot No.','EYA07')]:
            d.add_paragraph(label);d.add_paragraph(':');d.add_paragraph(text)
        t=d.add_table(rows=2,cols=5)
        for c,v in zip(t.rows[0].cells,['ANALYSIS DESIRED:','Area','Standard Specifications','Results','Remarks']):c.text=v
        for c,v in zip(t.rows[1].cells,['Standard Plate Count (SPC)','Mixing tank','Nmt 100 cfu/mL',value,'Passed']):c.text=v
        d.save(file);return file.read_bytes()
    def test_archive_deduplicates_patterns_without_importing_results(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);archive=folder/'source.zip'
            a=self.report(folder,'Environmental Monitoring','98765321');b=self.report(folder,'Environmental Monitoring','44444444')
            with zipfile.ZipFile(archive,'w') as z:z.writestr('PF/2026/Boost/a.docx',a);z.writestr('PF/2026/Boost/b.docx',b)
            result=scan(archive,folder/'prepared');self.assertEqual(len(result['candidates']),1);self.assertEqual(len(result['candidates'][0]['evidence']),2)
            self.assertNotIn('98765321',json.dumps(result));self.assertNotIn('44444444',json.dumps(result));self.assertNotIn('Passed',json.dumps(result))
    def test_non_environmental_category_is_not_reclassified(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);archive=folder/'source.zip'
            with zipfile.ZipFile(archive,'w') as z:z.writestr('PF/2026/Boost/report.docx',self.report(folder,'Goods-in-Process','4'))
            result=scan(archive,folder/'prepared');self.assertEqual(result['candidates'],[]);self.assertIn('do not reclassify',result['exceptions'][0]['reason'])
    def test_flat_archive_uses_header_product_and_process_area(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);archive=folder/'source.zip'
            with zipfile.ZipFile(archive,'w') as z:z.writestr('report.docx',self.report(folder,'Environmental Monitoring','4'))
            result=scan(archive,folder/'prepared');candidate=result['candidates'][0]
            self.assertEqual(candidate['productHint'],'Boost');self.assertEqual(candidate['areaHint'],'Compounding')

if __name__=='__main__':unittest.main()
