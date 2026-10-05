from docx import Document
from zipfile import ZipFile
from lxml import etree as E
from pathlib import Path
files=[Path(r'C:\Users\Roy\Documents\ipi format\ENVI.docx'),Path(r'C:\Users\Roy\Documents\ChatGPT\IPI\output\Environmental Monitoring Tagged Template Candidate.docx')]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for f in files:
 print('\n===',f,'===')
 d=Document(f)
 print('sections',len(d.sections),'page_inches',[(round(s.page_width.inches,2),round(s.page_height.inches,2),round(s.top_margin.inches,2),round(s.bottom_margin.inches,2),round(s.left_margin.inches,2),round(s.right_margin.inches,2)) for s in d.sections])
 print('body paragraphs',[p.text for p in d.paragraphs if p.text.strip()])
 print('body tables',len(d.tables))
 for ti,t in enumerate(d.tables):
  print('TABLE',ti,'rows',len(t.rows),'grid_cols',len(t.columns))
  for ri,row in enumerate(t.rows[:20]):
   print(ri,' | '.join(c.text.replace('\n',' / ') for c in row.cells))
 with ZipFile(f) as z:
  for part in ('word/document.xml','word/header1.xml','word/footer1.xml'):
   if part not in z.namelist():continue
   root=E.fromstring(z.read(part))
   print(part,'vMerge',len(root.xpath('.//w:vMerge',namespaces=ns)),'gridSpan',len(root.xpath('.//w:gridSpan',namespaces=ns)),'drawings',len(root.xpath('.//w:drawing|.//w:pict',namespaces=ns)),'fields',[x.text for x in root.xpath('.//w:instrText',namespaces=ns) if x.text])
